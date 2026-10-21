"""Put the modelling figures into the presentation-video script, at the cue points.

The script document carries tracked changes, so the paragraph text python-docx
reports is not what the document says: deleted runs are still present and
inserted runs sit inside ``w:ins``. Anchors are therefore matched against the
*accepted* text -- every ``w:t`` under the paragraph, which excludes
``w:delText`` -- and nothing existing is modified, only new paragraphs added
between existing ones. Rejecting or accepting the revisions later leaves the
figures in place either way.

Re-running replaces the figures it inserted before, so it is safe after a
figure is redrawn or after the script is edited around them.

    PYTHONPATH=. .venv/bin/python insert_script_figures.py
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SCRIPT = ROOT / "FLAME_Presentation_Video_建模组回改_1016.docx"
BACKUP = ROOT / ".docx-build" / "FLAME_Presentation_Video_建模组回改_1016_无配图原稿.docx"
CUES = ROOT / "outputs" / "video_cues"

#: Caption paragraphs start with this, which is how a re-run finds its own work.
MARKER = "【配图"

IMAGE_WIDTH = Cm(16.4)
CAPTION_INK = RGBColor(0x52, 0x51, 0x4E)

#: (file, anchor in the English script, English caption, anchor in the Chinese
#: script, Chinese caption). Anchors are matched against accepted text.
PLAN: list[tuple[str, str, str, str, str]] = [
    (
        "01_circuit.png",
        "The hypoxia-responsive promoter pPepT controls LuxI",
        "the two sensing arms, and the AND gate they feed.",
        "低氧响应型启动子 pPepT 控制 LuxI",
        "两条传感臂，以及它们汇成的 AND 门。",
    ),
    (
        "02_and_gate_holdout.png",
        "First, it tested whether the AND gate acts as a true product",
        "hold one corner out and predict it from the other three: "
        "98.13 % predicted against 98.08 % measured at 24 h, while the same "
        "check fails at 9 h.",
        "第一，建模检验 AND 门是否真是两个输入的乘积",
        "扣掉一个角，用另外三个角把它预测出来：24 h 预测 98.13 %，实测 98.08 %；"
        "同一检验在 9 h 失败。",
    ),
    (
        "03_arms_anti_selective.png",
        "Second, modeling asked where else the circuit might activate",
        "each designed arm on its own is anti-selective in the lumen: "
        "below 1.0 it reports the off-target site more strongly than the tumour.",
        "第二，建模追问线路还可能在哪里开启",
        "单看每条设计臂，在肠腔都是反向选择：低于 1.0 即对非靶位点的响应强于肿瘤。",
    ),
    (
        "04_tissue_axis_bottleneck.png",
        "Second, modeling asked where else the circuit might activate",
        "with the native lldPRD repression included, the bottleneck "
        "moves from the lumen to the colonic mucosa, at 21.85×.",
        "第二，建模追问线路还可能在哪里开启",
        "计入天然 lldPRD 抑制后，瓶颈从肠腔移到结肠黏膜，为 21.85×。",
    ),
    (
        "05_lesion_floor_178um.png",
        "Modeling also identified a geometric safety floor",
        "oxygen reaches a fixed depth, so a lesion below 178 µm has "
        "no hypoxic volume at all. Pure geometry, no calibration.",
        "建模还给出一条几何安全底线",
        "氧只穿透固定深度，半径低于 178 µm 的病灶根本没有缺氧体积。纯几何，不需要标定。",
    ),
    (
        "06_dose_vs_margin.png",
        "First, modeling compared dose requirements with safety margins",
        "reaching a 150 µm lesion costs 4.95× dose and leaves 4.41× "
        "mucosa margin, under the 10× we require.",
        "第一，建模把所需剂量与安全余量放在一起比较",
        "够到 150 µm 病灶要 4.95× 剂量，黏膜余量只剩 4.41×，低于我们自己要求的 10×。",
    ),
    (
        "07_micrometastasis_cluster.png",
        "A second model, built on completely different data",
        "even 125 micro-lesions clustered together stay below the "
        "quorum threshold: not a problem a larger dose can solve.",
        "另一个建立在完全不同数据上的模型",
        "哪怕 125 个微转移灶聚在一起，仍低于群体感应阈值：加大剂量解决不了。",
    ),
    (
        "08_immune_feasibility.png",
        "Finally, modeling explored whether an immune response",
        "a feasibility boundary, not a prediction: below εG = 1 no "
        "pyroptotic flux is enough, and a maximum treatable burden exists.",
        "最后，建模考察了免疫反应能否把作用延伸到远处病灶",
        "给出的是可行性边界而非预测：低于 εG = 1 时无论焦亡通量多大都不够，"
        "且存在一个最大可治负荷。",
    ),
    (
        "09_contribution.png",
        "Across the second and third questions, modeling contributed",
        "measured data in, decisions out.",
        "围绕第二和第三个问题，建模贡献了",
        "实测数据进，项目决定出。",
    ),
]


def accepted_text(paragraph) -> str:
    """What the paragraph says once the revisions are accepted.

    ``w:t`` covers ordinary and inserted runs; deleted runs carry ``w:delText``
    and are left out, which is what makes this match the final wording.
    """
    return "".join(node.text or "" for node in paragraph._p.iter(qn("w:t")))


def has_picture(paragraph) -> bool:
    return paragraph._p.find(qn("w:r") + "/" + qn("w:drawing")) is not None or any(
        True for _ in paragraph._p.iter(qn("a:blip"))
    )


def drop_previous_figures(document) -> int:
    """Remove the figure blocks a previous run inserted."""
    removed = 0
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.lstrip().startswith(MARKER):
            continue
        if index and has_picture(paragraphs[index - 1]):
            node = paragraphs[index - 1]._p
            node.getparent().remove(node)
            removed += 1
        node = paragraph._p
        node.getparent().remove(node)
    return removed


def insert_after(document, anchor):
    """A new empty paragraph immediately after ``anchor``."""
    created = document.add_paragraph()
    anchor._p.addnext(created._p)
    return created


def find(document, needle: str, *, after: int = 0) -> int:
    for index, paragraph in enumerate(document.paragraphs):
        if index >= after and needle in accepted_text(paragraph):
            return index
    raise LookupError(needle)


def place(document, anchor_index: int, number: int, image: Path, caption: str):
    anchor = document.paragraphs[anchor_index]
    picture = insert_after(document, anchor)
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(8)
    picture.paragraph_format.space_after = Pt(2)
    picture.add_run().add_picture(str(image), width=IMAGE_WIDTH)

    legend = insert_after(document, picture)
    legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legend.paragraph_format.space_after = Pt(10)
    run = legend.add_run(f"{MARKER} {number}】{caption}")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = CAPTION_INK
    return legend


def build() -> Path:
    if not BACKUP.exists():
        BACKUP.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(SCRIPT, BACKUP)

    document = Document(str(SCRIPT))
    removed = drop_previous_figures(document)
    if removed:
        print(f"replaced {removed} figure(s) from an earlier run")

    # The Chinese script repeats every anchor, so each half is searched from the
    # heading that starts it rather than from the top of the document.
    chinese_start = find(document, "中文演讲稿")

    for number, (name, en_anchor, en_caption, zh_anchor, zh_caption) in enumerate(PLAN, 1):
        image = CUES / name
        if not image.exists():
            raise FileNotFoundError(f"{image} -- run build_video_cues.py first")
        for anchor, caption, after in (
            (en_anchor, en_caption, 0),
            (zh_anchor, zh_caption, chinese_start),
        ):
            index = find(document, anchor, after=after)
            # Later figures on the same anchor must land after the earlier ones.
            while index + 1 < len(document.paragraphs) and (
                has_picture(document.paragraphs[index + 1])
                or document.paragraphs[index + 1].text.lstrip().startswith(MARKER)
            ):
                index += 1
            place(document, index, number, image, caption)
            chinese_start = find(document, "中文演讲稿")

    document.save(str(SCRIPT))
    return SCRIPT


if __name__ == "__main__":
    path = build()
    size = path.stat().st_size / 1024
    print(f"{path.name}  ({size:.0f} KB)")
    print(f"original kept at {BACKUP.relative_to(ROOT)}")
