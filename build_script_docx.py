"""Build the Word version of the presentation-video script.

Single source: `PRESENTATION_VIDEO_MODELING.md`. The Markdown stays the thing
people edit; this only renders it.

The renderer is deliberately narrow — it understands exactly the constructs that
file uses (headings, the `> **EN** —` / `> **中** —` paired blockquotes, the
`**画面**：` cue lines, tables, lists, fenced blocks and inline bold/code/links)
and nothing else. A spoken script is read aloud off the page, so the English
voice-over blocks are set large with wide leading, and the Chinese gloss sits
under them in a smaller grey face.

    .venv/bin/python build_script_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PRESENTATION_VIDEO_MODELING.md"
OUT = ROOT / "建模讲稿_PresentationVideo.docx"

HAN = "方正书宋简体" if False else "宋体"
HAN_BOLD = "黑体"
LATIN = "Calibri"
MONO = "Consolas"

INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x1F, 0x5C, 0x3A)


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #

def style_run(run, *, size=10.5, bold=False, italic=False, color=INK,
              latin=LATIN, han=HAN, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = MONO if mono else latin
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), MONO if mono else latin)
    fonts.set(qn("w:hAnsi"), MONO if mono else latin)
    fonts.set(qn("w:eastAsia"), MONO if mono else han)
    return run


def shade(paragraph, hexcolor):
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pr.append(shd)


def left_bar(paragraph, hexcolor, size=18):
    pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hexcolor)
    borders.append(left)
    pr.append(borders)


def keep_with_next(paragraph):
    pr = paragraph._p.get_or_add_pPr()
    pr.append(OxmlElement("w:keepNext"))


INLINE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\n]+\*)"
)


def emit_inline(paragraph, text, **base):
    """Write `text` into `paragraph`, honouring bold / code / link / italic."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        kw = dict(base)
        if piece.startswith("**") and piece.endswith("**"):
            kw["bold"] = True
            kw["han"] = HAN_BOLD
            piece = piece[2:-2]
        elif piece.startswith("`") and piece.endswith("`"):
            kw["mono"] = True
            kw["size"] = kw.get("size", 10.5) - 0.5
            piece = piece[1:-1]
        elif piece.startswith("[") and "](" in piece:
            label, url = piece[1:-1].split("](", 1)
            kw["color"] = ACCENT
            piece = label if url.startswith("#") else f"{label}"
        elif piece.startswith("*") and piece.endswith("*"):
            kw["italic"] = True
            piece = piece[1:-1]
        style_run(paragraph.add_run(piece), **kw)


def strip_md(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


# --------------------------------------------------------------------------- #
# block writers
# --------------------------------------------------------------------------- #

def add_para(doc, *, space_before=0, space_after=6, left=0, line=1.3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.left_indent = Cm(left)
    pf.line_spacing = line
    return p


def heading(doc, level, text):
    sizes = {1: 20, 2: 15, 3: 12.5}
    befores = {1: 18, 2: 16, 3: 12}
    p = add_para(doc, space_before=befores[level], space_after=6, line=1.15)
    keep_with_next(p)
    emit_inline(p, text, size=sizes[level], bold=True, han=HAN_BOLD,
                color=INK if level < 3 else ACCENT)
    if level == 2:
        pr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "C8C8C8")
        bd.append(bottom)
        pr.append(bd)


def vo_block(doc, text, *, english):
    """The spoken lines. English is the thing actually read aloud."""
    body = text.split("—", 1)[1].strip() if "—" in text.split("\n")[0] else text
    p = add_para(doc, space_before=4, space_after=4, left=0.3,
                 line=1.5 if english else 1.3)
    left_bar(p, "1F5C3A" if english else "C8C8C8")
    shade(p, "F4F7F4" if english else "FAFAFA")
    tag = p.add_run("EN　" if english else "中　")
    style_run(tag, size=9.5, bold=True, han=HAN_BOLD,
              color=ACCENT if english else GREY)
    emit_inline(p, body, size=12 if english else 10,
                color=INK if english else GREY)


def cue_block(doc, text):
    p = add_para(doc, space_before=4, space_after=4, left=0.0, line=1.2)
    shade(p, "F2F2F2")
    emit_inline(p, text, size=9.5, color=GREY)


def note_block(doc, text):
    p = add_para(doc, space_before=4, space_after=6, left=0.3, line=1.3)
    left_bar(p, "B0B0B0", size=12)
    emit_inline(p, text, size=10, color=GREY)


def code_block(doc, lines):
    for line in lines:
        p = add_para(doc, space_after=0, left=0.3, line=1.1)
        shade(p, "F5F5F5")
        style_run(p.add_run(line or " "), size=9, mono=True, color=GREY)


def bullet(doc, text, ordered=False, idx=1):
    p = add_para(doc, space_after=3, left=0.6, line=1.25)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    marker = f"{idx}. " if ordered else "• "
    style_run(p.add_run(marker), size=10.5, color=GREY)
    emit_inline(p, text, size=10.5)


def table_block(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, cell in enumerate(header):
        para = t.cell(0, j).paragraphs[0]
        para.paragraph_format.space_after = Pt(2)
        emit_inline(para, cell, size=9.5, bold=True, han=HAN_BOLD)
        shade(para, "EDF2ED")
    for i, row in enumerate(body, start=1):
        for j, cell in enumerate(row):
            if j >= len(header):
                continue
            para = t.cell(i, j).paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.15
            emit_inline(para, cell, size=9.5)
    add_para(doc, space_after=4)


def rule(doc):
    p = add_para(doc, space_before=6, space_after=6)
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D8D8D8")
    bd.append(bottom)
    pr.append(bd)


# --------------------------------------------------------------------------- #
# document frame
# --------------------------------------------------------------------------- #

def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), HAN)

    # footer: page numbers
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    style_run(run, size=9, color=GREY)
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._element.append(el)
    return doc


def title_page(doc, lines):
    p = add_para(doc, space_before=24, space_after=2)
    style_run(p.add_run("Presentation Video — Modeling"),
              size=24, bold=True, han=HAN_BOLD)
    p = add_para(doc, space_after=14)
    style_run(p.add_run("建模部分口播稿"), size=15, bold=True,
              han=HAN_BOLD, color=ACCENT)
    for line in lines:
        p = add_para(doc, space_after=3, left=0.0, line=1.35)
        emit_inline(p, line, size=10, color=GREY)
    rule(doc)


# --------------------------------------------------------------------------- #
# parser
# --------------------------------------------------------------------------- #

def convert(md: str, doc: Document) -> None:
    lines = md.split("\n")
    i = 0
    pending_quote: list[str] = []
    header_done = False
    ordered_idx = 0

    def flush_quote():
        nonlocal pending_quote
        if not pending_quote:
            return
        text = " ".join(x.strip() for x in pending_quote).strip()
        pending_quote = []
        if text.startswith("**EN**"):
            vo_block(doc, text, english=True)
        elif text.startswith("**中**"):
            vo_block(doc, text, english=False)
        else:
            note_block(doc, text)

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith(">"):
            body = line[1:].lstrip()
            if not body and pending_quote:
                flush_quote()
            elif body.startswith("**EN**") or body.startswith("**中**"):
                flush_quote()
                pending_quote.append(body)
            else:
                pending_quote.append(body)
            i += 1
            continue
        flush_quote()

        if not line.strip():
            i += 1
            continue

        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, block)
            i += 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                table_block(doc, rows)
            continue

        if re.fullmatch(r"-{3,}", line.strip()):
            rule(doc)
            i += 1
            continue

        if line.startswith("### "):
            heading(doc, 3, line[4:])
            i += 1
            continue
        if line.startswith("## "):
            heading(doc, 2, line[3:])
            i += 1
            continue
        if line.startswith("# "):
            if not header_done:
                # title block: the H1 plus the leading quoted intro
                intro = []
                j = i + 1
                while j < len(lines) and not lines[j].startswith("## "):
                    s = lines[j].lstrip("> ").strip()
                    if s and not re.fullmatch(r"-{3,}", s):
                        intro.append(s)
                    j += 1
                title_page(doc, intro)
                header_done = True
                i = j
                continue
            heading(doc, 1, line[2:])
            i += 1
            continue

        stripped = line.strip()
        m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m:
            ordered_idx += 1
            bullet(doc, m.group(2), ordered=True, idx=int(m.group(1)))
            i += 1
            continue
        ordered_idx = 0

        if stripped.startswith(("- ", "* ")):
            bullet(doc, stripped[2:])
            i += 1
            continue

        if stripped.startswith("**画面**"):
            cue_block(doc, stripped)
            i += 1
            continue

        p = add_para(doc, space_after=6, line=1.35)
        emit_inline(p, stripped, size=10.5)
        i += 1

    flush_quote()


def main() -> int:
    if not SOURCE.exists():
        print(f"missing source: {SOURCE}", file=sys.stderr)
        return 1
    doc = make_doc()
    convert(SOURCE.read_text(encoding="utf-8"), doc)
    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({OUT.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
